#!/usr/bin/env python3
"""Generate the long-form ChildGuard presentation (43 slides)."""
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / "docs" / "assets"
OUTPUT = ROOT / "docs" / "ChildGuard_3_Tuan.pptx"

NAVY = RGBColor(10, 18, 34)
PANEL = RGBColor(29, 43, 66)
PANEL_2 = RGBColor(36, 55, 82)
WHITE = RGBColor(245, 248, 255)
MUTED = RGBColor(160, 177, 204)
BLUE = RGBColor(59, 130, 246)
CYAN = RGBColor(34, 211, 238)
GREEN = RGBColor(16, 185, 129)
ORANGE = RGBColor(245, 158, 11)
RED = RGBColor(239, 68, 68)


def add_text(slide, text, x, y, w, h, size=22, color=WHITE, bold=False,
             align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def add_bullets(slide, items, x, y, w, h, size=21, color=WHITE, accent=CYAN):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(10)
        p.text = "•  " + p.text
    return box


def rect(slide, x, y, w, h, fill=PANEL, radius=True, line=None):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line or fill
    return shp


def base_slide(prs, title, section="", accent=BLUE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = NAVY
    rect(slide, 0, 0, 13.333, 0.08, accent, False)
    add_text(slide, title, 0.65, 0.34, 11.7, 0.55, 25, WHITE, True)
    if section:
        add_text(slide, section.upper(), 10.75, 0.44, 1.9, 0.28, 10, accent, True, PP_ALIGN.RIGHT)
    return slide


def footer(slide, number, section):
    add_text(slide, "ChildGuard · An Open Guard Kids Application", 0.65, 7.16, 7.2, 0.22, 9, MUTED)
    add_text(slide, f"{section}   ·   {number:02d}", 10.2, 7.16, 2.45, 0.22, 9, MUTED, False, PP_ALIGN.RIGHT)


def picture_contain(slide, path, x, y, w, h, padding=0.08):
    path = Path(path)
    with Image.open(path) as im:
        iw, ih = im.size
    ratio = min((w - 2*padding) / iw, (h - 2*padding) / ih)
    pw, ph = iw * ratio, ih * ratio
    px, py = x + (w-pw)/2, y + (h-ph)/2
    slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = NAVY
    rect(slide, 0, 0, 13.333, 0.12, CYAN, False)
    rect(slide, 0.72, 0.72, 1.1, 1.1, BLUE)
    add_text(slide, "CG", 0.72, 0.79, 1.1, 0.8, 30, WHITE, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, "CHILDGUARD", 2.05, 0.83, 6.8, 0.5, 18, CYAN, True)
    add_text(slide, "AN OPEN GUARD KIDS APPLICATION", 2.05, 1.24, 7.0, 0.35, 12, MUTED, True)
    add_text(slide, "Hành trình xây dựng hệ thống\nqua 3 tuần phát triển", 0.8, 2.05, 8.0, 1.55, 31, WHITE, True)
    add_text(slide, "Agent Windows · FastAPI Server · Dashboard phụ huynh", 0.83, 3.78, 8.7, 0.45, 17, MUTED)
    rect(slide, 9.55, 1.75, 2.65, 3.65, PANEL)
    add_text(slide, "03", 9.75, 2.15, 2.2, 1.1, 54, CYAN, True, PP_ALIGN.CENTER)
    add_text(slide, "TUẦN PHÁT TRIỂN", 9.7, 3.25, 2.3, 0.35, 12, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, "Kiến trúc\nThời gian\nNội dung & báo cáo", 9.85, 3.82, 2.0, 1.1, 15, MUTED, False, PP_ALIGN.CENTER)
    add_text(slide, "Nguyễn Văn Phước — 23120074\nHà Công Thuận — 23120094", 0.83, 5.35, 6.2, 0.8, 16, WHITE, True)
    add_text(slide, "GVHD: Đỗ Hoàng Cường\nTrường Đại học Khoa học Tự nhiên, ĐHQG-HCM", 0.83, 6.25, 8.0, 0.62, 12, MUTED)
    footer(slide, 1, "Mở đầu")


def section_slide(prs, number, label, title, subtitle, accent):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = NAVY
    rect(slide, 0, 0, 13.333, 7.5, accent, False)
    rect(slide, 0.15, 0.15, 13.033, 7.2, NAVY, False)
    add_text(slide, label.upper(), 0.9, 1.15, 4.0, 0.4, 15, accent, True)
    add_text(slide, title, 0.9, 2.0, 10.8, 1.45, 34, WHITE, True)
    add_text(slide, subtitle, 0.92, 3.68, 10.5, 0.9, 18, MUTED)
    add_text(slide, f"{number:02d}", 10.2, 4.9, 2.0, 1.15, 58, accent, True, PP_ALIGN.RIGHT)
    footer(slide, number, label)


def bullet_slide(prs, number, title, items, section, accent=BLUE, note=None):
    slide = base_slide(prs, title, section, accent)
    rect(slide, 0.65, 1.15, 12.0, 5.65, PANEL)
    add_bullets(slide, items, 1.0, 1.48, 11.3, 4.85, 20 if len(items) <= 6 else 17)
    if note:
        rect(slide, 1.0, 5.95, 11.3, 0.55, PANEL_2)
        add_text(slide, note, 1.2, 6.08, 10.9, 0.25, 12, accent, True, PP_ALIGN.CENTER)
    footer(slide, number, section)


def image_slide(prs, number, title, path, section, caption, accent=BLUE):
    slide = base_slide(prs, title, section, accent)
    rect(slide, 0.65, 1.1, 12.0, 5.75, PANEL)
    picture_contain(slide, path, 0.85, 1.28, 11.6, 4.95)
    add_text(slide, caption, 0.9, 6.25, 11.5, 0.35, 11, MUTED, False, PP_ALIGN.CENTER)
    footer(slide, number, section)


def code_slide(prs, number, title, code, bullets, section, accent=BLUE):
    slide = base_slide(prs, title, section, accent)
    rect(slide, 0.65, 1.12, 6.25, 5.7, RGBColor(5, 11, 22))
    add_text(slide, code, 0.95, 1.45, 5.65, 5.05, 15, RGBColor(167, 243, 208), False, font="Consolas")
    rect(slide, 7.15, 1.12, 5.5, 5.7, PANEL)
    add_bullets(slide, bullets, 7.5, 1.5, 4.85, 4.95, 18)
    footer(slide, number, section)


def table_slide(prs, number, title, headers, rows, section, widths=None, accent=BLUE, font_size=14):
    slide = base_slide(prs, title, section, accent)
    x, y, w, h = 0.65, 1.2, 12.0, 5.55
    shape = slide.shapes.add_table(len(rows)+1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    table = shape.table
    if widths:
        for idx, width in enumerate(widths): table.columns[idx].width = Inches(width)
    for c, header in enumerate(headers):
        cell = table.cell(0, c); cell.fill.solid(); cell.fill.fore_color.rgb = accent
        cell.text = header
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row):
            cell = table.cell(r, c); cell.fill.solid(); cell.fill.fore_color.rgb = PANEL if r % 2 else PANEL_2
            cell.text = str(val)
    for r in range(len(rows)+1):
        for c in range(len(headers)):
            cell = table.cell(r, c)
            cell.margin_left = Inches(.08); cell.margin_right = Inches(.08)
            cell.margin_top = Inches(.05); cell.margin_bottom = Inches(.05)
            for p in cell.text_frame.paragraphs:
                p.font.name = "Aptos"; p.font.size = Pt(font_size)
                p.font.bold = r == 0; p.font.color.rgb = WHITE
                p.alignment = PP_ALIGN.CENTER if r == 0 else PP_ALIGN.LEFT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    footer(slide, number, section)


def flow_slide(prs, number, title, steps, section, accent=BLUE, note=None):
    slide = base_slide(prs, title, section, accent)
    count = len(steps); gap = 0.16; total_w = 12.0; box_w = (total_w-gap*(count-1))/count
    for i, (head, body) in enumerate(steps):
        x = 0.65 + i*(box_w+gap)
        rect(slide, x, 2.0, box_w, 2.8, PANEL if i % 2 == 0 else PANEL_2)
        add_text(slide, f"{i+1:02d}", x+0.18, 2.2, box_w-0.36, 0.45, 22, accent, True)
        add_text(slide, head, x+0.18, 2.82, box_w-0.36, 0.5, 16, WHITE, True)
        add_text(slide, body, x+0.18, 3.45, box_w-0.36, 1.0, 12, MUTED)
        if i < count-1:
            add_text(slide, "→", x+box_w-0.03, 3.0, gap+0.1, 0.45, 20, accent, True, PP_ALIGN.CENTER)
    if note:
        add_text(slide, note, 0.85, 5.35, 11.6, 0.65, 15, accent, True, PP_ALIGN.CENTER)
    footer(slide, number, section)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    title_slide(prs)
    bullet_slide(prs, 2, "Bài toán cần giải quyết", [
        "Quản lý thời gian sử dụng máy tính mà không cần phụ huynh luôn ở cạnh.",
        "Hạn chế ứng dụng và website không phù hợp theo chính sách gia đình.",
        "Duy trì chức năng cốt lõi khi mạng hoặc máy chủ bị gián đoạn.",
        "Giải thích rõ cho trẻ vì sao một hành động bị chặn.",
        "Chỉ thu thập metadata cần thiết; không đọc nội dung riêng tư."
    ], "Mở đầu", CYAN, "Công cụ đồng hành — không phải phần mềm theo dõi lén")
    table_slide(prs, 3, "Lộ trình phát triển", ["Tuần 1", "Tuần 2", "Tuần 3"], [
        ["Kiến trúc", "Quota và lịch", "DNS và SafeSearch"],
        ["Mô hình dữ liệu", "Monotonic clock", "Chặn ứng dụng"],
        ["Enrollment", "Đồng bộ policy", "Event và báo cáo"],
        ["Heartbeat & Dashboard", "Khóa máy", "Giao diện cho trẻ"]
    ], "Mở đầu", [4,4,4], CYAN, 16)

    section_slide(prs, 4, "Tuần 1", "Xây dựng nền tảng hệ thống", "Kiến trúc · Dữ liệu · Enrollment · Heartbeat · Dashboard", BLUE)
    image_slide(prs, 5, "Kiến trúc tổng thể", ASSETS/"architecture.png", "Tuần 1", "Agent Windows ↔ FastAPI Server ↔ Dashboard phụ huynh", BLUE)
    table_slide(prs, 6, "Vai trò của từng thành phần", ["Thành phần", "Trách nhiệm"], [
        ["Agent Core", "Quota, lịch, chặn ứng dụng và website"], ["Tray UI", "Trạng thái, cảnh báo, yêu cầu thêm giờ"],
        ["Sync Worker", "Heartbeat, policy, event và command"], ["FastAPI", "Xác thực, thiết bị, policy, audit"],
        ["Dashboard", "Giao diện quản lý của phụ huynh"], ["SQLite", "Cache cục bộ và dữ liệu server"]
    ], "Tuần 1", [3.2,8.8], BLUE, 15)
    code_slide(prs, 7, "Cấu trúc mã nguồn", """app/\n├── agent/\n│   ├── enforcer.py\n│   ├── storage.py\n│   ├── sync.py\n│   ├── dns_proxy.py\n│   └── tray/\n├── server/app/\n│   ├── routers/\n│   ├── models.py\n│   └── schemas.py\n├── dashboard/\n└── tests/""", ["Agent chịu trách nhiệm thực thi", "Server quản lý và đồng bộ", "Dashboard không điều khiển trực tiếp", "Giao tiếp qua API rõ ràng"], "Tuần 1", BLUE)
    image_slide(prs, 8, "Mô hình dữ liệu", ASSETS/"data_model.png", "Tuần 1", "10 thực thể chính; Policy được lưu theo phiên bản", BLUE)
    bullet_slide(prs, 9, "Vì sao chọn SQLite và SQLAlchemy?", [
        "Cài nhanh, không cần database server riêng.", "SQLite WAL phù hợp quy mô phòng lab.",
        "SQLAlchemy tách mô hình dữ liệu khỏi SQL cụ thể.", "Dễ tạo database tạm cho kiểm thử tự động.",
        "Giới hạn: khi mở rộng nên chuyển PostgreSQL và migration chuẩn."
    ], "Tuần 1", BLUE)
    image_slide(prs, 10, "Luồng ghép đôi thiết bị", ASSETS/"enrollment.png", "Tuần 1", "Mã 8 ký tự · hạn 10 phút · token được phát sau khi xác minh", BLUE)
    image_slide(prs, 11, "Giao diện tạo mã ghép đôi", ASSETS/"preview/enroll.png", "Tuần 1", "Ảnh preview phiên bản cũ — cần chụp lại trước khi bảo vệ", BLUE)
    flow_slide(prs, 12, "Access token và refresh token", [
        ("Enrollment", "Mã ghép đôi hợp lệ"), ("Access token", "Hiện tại: 10 phút"),
        ("Refresh token", "Hiệu lực 30 ngày"), ("Token hash", "Server chỉ lưu SHA-256")
    ], "Tuần 1", BLUE, "Token được gắn với device_id và bị thu hồi khi thiết bị không còn hoạt động")
    code_slide(prs, 13, "Cơ chế heartbeat", """POST /heartbeat\n{\n  \"policy_version\": 3,\n  \"quota_used_seconds\": 5400,\n  \"queued_event_count\": 2\n}\n\nResponse:\n{\n  \"policy_changed\": true,\n  \"pending_commands\": [],\n  \"server_time\": \"...Z\"\n}""", ["Cập nhật last_seen", "So sánh policy version", "Đồng bộ thời gian UTC", "Nhận command dự phòng", "Theo dõi event queue"], "Tuần 1", BLUE)
    image_slide(prs, 14, "Trạng thái thiết bị trên Dashboard", ASSETS/"preview/overview-dashboard.png", "Tuần 1", "Thiết bị online/offline · quota · policy version · điều khiển khẩn", BLUE)

    section_slide(prs, 15, "Tuần 2", "Quản lý thời gian và chính sách", "Quota · Lịch 7×48 · Cảnh báo · Khóa máy · Policy sync", GREEN)
    code_slide(prs, 16, "Cấu trúc một policy", """{\n  \"version\": 15,\n  \"quota\": {\n    \"daily_minutes\": 90\n  },\n  \"schedule\": {\n    \"matrix\": \"7 x 48\"\n  },\n  \"app_rules\": {...},\n  \"domain_rules\": {...}\n}""", ["Version tăng dần", "Agent cache trong SQLite", "Áp dụng được khi offline", "Chưa có HMAC/chữ ký toàn vẹn"], "Tuần 2", GREEN)
    code_slide(prs, 17, "Thuật toán quota", """Mỗi tick:\n  now = monotonic()\n  delta = now - previous\n\n  nếu không khóa\n  và idle ≤ 5 phút:\n      quota += delta\n\n  nếu quota hết:\n      cảnh báo / ân hạn\n      khóa sau 60 giây""", ["Chỉ tính thời gian sử dụng thực", "Lưu số giây trong SQLite", "Không reset khi Agent restart", "Cần bổ sung quota riêng ngày thường/cuối tuần"], "Tuần 2", GREEN)
    table_slide(prs, 18, "Vì sao dùng monotonic clock?", ["Nguồn thời gian", "Vai trò", "Rủi ro"], [
        ["time.time()", "Thời gian hệ thống", "Có thể bị chỉnh lùi"],
        ["time.monotonic()", "Đo khoảng thời gian", "Chỉ tăng trong một phiên chạy"],
        ["Server timestamp", "Xác định UTC/ngày", "Phụ thuộc lần đồng bộ gần nhất"]
    ], "Tuần 2", [2.5,4.5,5], GREEN, 15)
    bullet_slide(prs, 19, "Phát hiện idle và màn hình khóa", [
        "GetLastInputInfo tính thời gian không có thao tác bàn phím/chuột.",
        "Không cộng quota khi idle quá 300 giây.", "Phát hiện lock screen bằng trạng thái desktop Windows.",
        "Tray UI vẫn hiển thị trạng thái trong phiên người dùng.", "Các API Windows cần được nghiệm thu trên máy thật."
    ], "Tuần 2", GREEN)
    table_slide(prs, 20, "Lịch tuần 7×48", ["Thành phần", "Ý nghĩa"], [
        ["7 hàng", "Thứ Hai đến Chủ Nhật"], ["48 cột", "Mỗi ô tương ứng 30 phút"],
        ["Giá trị 1", "Được phép sử dụng"], ["Giá trị 0", "Ngoài khung giờ cho phép"],
        ["Tiêu chí", "Khóa trong 5 giây khi ngoài lịch"]
    ], "Tuần 2", [3,9], GREEN, 16)
    flow_slide(prs, 21, "Cảnh báo và thời gian ân hạn", [
        ("10 phút", "Cảnh báo lần 1"), ("5 phút", "Cảnh báo lần 2"), ("1 phút", "Cảnh báo khẩn"),
        ("Hết giờ", "Ân hạn 60 giây"), ("Khóa", "Ghi event locked")
    ], "Tuần 2", GREEN, "Không khóa đột ngột — trẻ có thời gian lưu bài")
    image_slide(prs, 22, "Giao diện cấu hình chính sách", ASSETS/"preview/policy.png", "Tuần 2", "Ảnh preview: quota · khung giờ · blacklist ứng dụng · blacklist tên miền", GREEN)
    flow_slide(prs, 23, "Luồng đồng bộ policy", [
        ("Phụ huynh", "Lưu chính sách"), ("Server", "Tạo version mới"), ("Heartbeat", "policy_changed=true"),
        ("Agent", "GET /policy"), ("Enforcer", "Cache và áp dụng")
    ], "Tuần 2", GREEN, "Mục tiêu đề bài: chính sách có hiệu lực trong vòng 60 giây")
    flow_slide(prs, 24, "Command thời gian thực", [
        ("Dashboard", "lock / unlock / grant"), ("Server", "Lưu Command"),
        ("WebSocket", "Đẩy tức thời"), ("Agent", "Thực thi"), ("ACK", "Xác nhận kết quả")
    ], "Tuần 2", GREEN, "Heartbeat là kênh dự phòng khi WebSocket gián đoạn")
    image_slide(prs, 25, "Demo khóa máy đúng hạn", ASSETS/"preview/test-lock.png", "Tuần 2", "Preview thử nghiệm Dashboard phát lệnh và máy Windows chuyển sang lock screen", GREEN)

    section_slide(prs, 26, "Tuần 3", "Kiểm soát nội dung và báo cáo", "DNS Proxy · SafeSearch · App Control · Event · Báo cáo · Minh bạch", ORANGE)
    flow_slide(prs, 27, "Kiến trúc lọc website", [
        ("Trình duyệt", "Phát DNS query"), ("Windows", "Stub resolver"), ("DNS Proxy", "127.0.0.1:53"),
        ("Policy", "Bypass / block / SafeSearch"), ("Phản hồi", "Upstream hoặc chặn")
    ], "Tuần 3", ORANGE, "Mỗi quyết định chặn tạo event và thông báo có lý do")
    code_slide(prs, 28, "Chuẩn hóa tên miền", """Đầu vào:\nhttps://www.news.example.co.uk\n:8443/path?q=abc\n\nChuẩn hóa:\nexample.co.uk\n\nLoại bỏ:\n- scheme / port\n- path / query\n- fragment""", ["Dùng Public Suffix List", "Chỉ lưu eTLD+1", "Không lưu URL đầy đủ", "Giảm dữ liệu riêng tư"], "Tuần 3", ORANGE)
    bullet_slide(prs, 29, "SafeSearch và danh sách bypass", [
        "Ép SafeSearch cho Google Search và Bing Search.", "Bật chế độ hạn chế phù hợp cho YouTube.",
        "Bypass cố định cho hạ tầng hệ thống và cổng giáo dục cần thiết.",
        "Không cho policy tự chặn dịch vụ khẩn cấp và bảo vệ trẻ em.",
        "Kiểm chứng bằng nslookup và trình duyệt trên Windows."
    ], "Tuần 3", ORANGE)
    bullet_slide(prs, 30, "Giới hạn của DNS filtering", [
        "DNS-over-HTTPS hoặc DNS-over-TLS có thể tạo đường vòng.", "VPN/proxy công cộng thay đổi đường đi lưu lượng.",
        "Truy cập trực tiếp bằng IP không có domain để so khớp.", "DNS không lọc được đường dẫn URL trong HTTPS.",
        "Người có quyền Administrator có thể đổi cấu hình mạng.", "Cần nêu biện pháp đối phó và rủi ro chấp nhận."
    ], "Tuần 3", ORANGE)
    flow_slide(prs, 31, "Chặn ứng dụng bằng tên và SHA-256", [
        ("psutil", "Liệt kê process"), ("Executable", "Tên + đường dẫn"), ("SHA-256", "Tính hoặc cache"),
        ("Policy", "So blacklist"), ("Enforcer", "Kill + notify + event")
    ], "Tuần 3", ORANGE, "Đổi tên executable không thay đổi SHA-256 của tệp")
    code_slide(prs, 32, "Hợp đồng sự kiện", """{\n  \"idempotency_key\": \"ev_6c47...\",\n  \"type\": \"blocked_domain\",\n  \"subject\": \"youtube.com\",\n  \"duration_sec\": 0,\n  \"policy_version\": 15,\n  \"created_at\": \"...Z\"\n}\n\n9 loại sự kiện hợp lệ""", ["Enum đóng", "Subject được chuẩn hóa", "Không chứa URL đầy đủ", "Server từ chối trường thừa", "Khác nhẹ schema đề bài: ts/policy_id"], "Tuần 3", ORANGE)
    flow_slide(prs, 33, "Đồng bộ sự kiện khi mất mạng", [
        ("Phát sinh", "Tạo event"), ("SQLite", "event_queue"), ("Batch", "Tối đa 100 mục"),
        ("Server", "Dedupe theo key"), ("Agent", "Xóa khi thành công")
    ], "Tuần 3", ORANGE, "At-least-once ở Agent · chống ghi trùng ở Server")
    table_slide(prs, 34, "Hạ tầng phân loại website", ["Nhóm", "Số lượng", "Mục đích"], [
        ["Giáo dục", "60", "Kiểm tra dương tính giả"], ["Giải trí", "40", "Phân loại báo cáo"],
        ["Mạng xã hội", "30", "Policy 6–11 tuổi"], ["Trò chơi", "30", "Theo dõi/kiểm soát"],
        ["Không phù hợp – mô phỏng", "20", "Không dùng nội dung người lớn thật"], ["Chưa phân loại", "20", "Fallback an toàn"]
    ], "Tuần 3", [4,2,6], ORANGE, 14)
    bullet_slide(prs, 35, "Báo cáo dành cho phụ huynh", [
        "Thời lượng sử dụng theo ngày và theo tuần.", "Top ứng dụng và top tên miền.",
        "Số lần ứng dụng hoặc website bị chặn.", "Phân loại hoạt động và sự kiện gần đây.",
        "Danh sách yêu cầu của trẻ.", "Xuất báo cáo tuần ra PDF."
    ], "Tuần 3", ORANGE)
    image_slide(prs, 36, "Giao diện yêu cầu từ trẻ", ASSETS/"preview/child-requests.png", "Tuần 3", "Preview phiên bản cũ — bản cuối nên có request pending và approved", ORANGE)
    bullet_slide(prs, 37, "Màn hình minh bạch dành cho trẻ", [
        "Hiển thị thời gian đã dùng và thời gian còn lại.", "Cho trẻ xem chính sách đang áp dụng.",
        "Thông báo ứng dụng/website bị chặn và lý do.", "Nút xin thêm 15/30/60 phút.",
        "Màn hình “Hôm nay em đã dùng gì”.", "Không có chế độ ẩn hoặc thu thập nội dung riêng tư."
    ], "Tuần 3", ORANGE, "NT1–NT5: trẻ biết, hiểu và có tiếng nói")
    flow_slide(prs, 38, "Xóa dữ liệu và thời hạn lưu", [
        ("90 ngày", "Retention tự động"), ("Phụ huynh", "Yêu cầu purge"), ("Server", "Xóa event/request"),
        ("Command", "clear_data"), ("Agent", "Xóa queue/bộ đếm")
    ], "Tuần 3", ORANGE, "Cần kiểm chứng dữ liệu bị xóa trên cả hai SQLite database")

    table_slide(prs, 39, "Kết quả sau ba tuần", ["Tuần", "Kết quả chính"], [
        ["Tuần 1", "Kiến trúc, database, enrollment, heartbeat và Dashboard"],
        ["Tuần 2", "Quota, lịch tuần, monotonic clock, policy và WebSocket"],
        ["Tuần 3", "DNS, SafeSearch, app control, event, báo cáo và giao diện trẻ"]
    ], "Tổng kết", [2.2,9.8], CYAN, 16)
    table_slide(prs, 40, "Mức độ hoàn thành F1–F5", ["Nhóm", "Đánh giá trung thực"], [
        ["F1", "Phần lớn logic đã có; thiếu quota riêng ngày thường/cuối tuần"],
        ["F2", "Blacklist/hash/DNS/SafeSearch có; DoH và allowlist còn thiếu"],
        ["F3", "Event/report/retention có; thiếu PRIVACY.md và cần thống nhất schema"],
        ["F4", "Version/WebSocket có; thiếu HMAC và audit old/new/IP"],
        ["F5", "Một phần; cần sửa auth, CSRF, ownership, encryption và STRIDE"]
    ], "Tổng kết", [2,10], CYAN, 14)
    bullet_slide(prs, 41, "Hạn chế và công việc cần hoàn thiện", [
        "Argon2id/scrypt, secret từ môi trường, Secure cookie và CSRF.",
        "Xác thực và ownership trên mọi API phụ huynh.", "HMAC/chữ ký policy; audit có IP và giá trị cũ/mới.",
        "Windows Service, ACL và bộ cài có recovery.", "Điều tra test suite đang treo và chạy lại toàn bộ.",
        "Đo CPU/RAM, DNS p50/p95 và latency với ≥30 mẫu."
    ], "Tổng kết", RED)
    table_slide(prs, 42, "Phân công và đóng góp", ["Thành viên", "Hạng mục minh chứng", "Tỷ lệ"], [
        ["Nguyễn Văn Phước\n23120074", "[Module · commit · test · tài liệu]", "[x%]"],
        ["Hà Công Thuận\n23120094", "[Module · commit · test · tài liệu]", "[100−x%]"]
    ], "Tổng kết", [3.3,6.8,1.9], CYAN, 16)
    slide = base_slide(prs, "Kết luận", "Q&A", CYAN)
    add_text(slide, "ChildGuard đã hình thành luồng end-to-end\ntừ phụ huynh cấu hình đến Agent thực thi\nvà gửi kết quả trở lại Dashboard.", 1.0, 1.5, 11.3, 1.65, 29, WHITE, True, PP_ALIGN.CENTER)
    rect(slide, 2.25, 3.65, 8.85, 1.15, PANEL)
    add_text(slide, "Minh bạch · Hoạt động ngoại tuyến · Tối thiểu dữ liệu", 2.45, 4.0, 8.45, 0.35, 18, CYAN, True, PP_ALIGN.CENTER)
    add_text(slide, "CẢM ƠN THẦY VÀ CÁC BẠN — Q&A", 2.2, 5.55, 8.95, 0.55, 22, WHITE, True, PP_ALIGN.CENTER)
    footer(slide, 43, "Kết thúc")

    prs.core_properties.title = "ChildGuard – Hành trình phát triển 3 tuần"
    prs.core_properties.subject = "Kiến trúc, quản lý thời gian, kiểm soát nội dung và báo cáo"
    prs.core_properties.author = "Nguyễn Văn Phước; Hà Công Thuận"
    prs.core_properties.comments = "Các ảnh giao diện được ghi rõ là preview phiên bản cũ."
    prs.save(OUTPUT)
    print(f"{OUTPUT}\nSlides: {len(prs.slides)}")


if __name__ == "__main__":
    build()
