#!/usr/bin/env python3
"""Render the editable Markdown report to a formatted Word document."""
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs" / "BaoCaoChinh_ChildGuard.md"
OUTPUT = ROOT / "docs" / "BaoCao.docx"
LEGACY_OUTPUT = ROOT / "docs" / "BaoCaoChinh_ChildGuard.docx"


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text.strip())
    r.bold = bold
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    separate.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def style_document(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(1)

    for name, size, before, after in (
        ("Title", 18, 12, 12),
        ("Heading 1", 15, 12, 8),
        ("Heading 2", 13, 10, 6),
        ("Heading 3", 12, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Report Note" not in doc.styles:
        note = doc.styles.add_style("Report Note", WD_STYLE_TYPE.PARAGRAPH)
        note.font.name = "Times New Roman"
        note.font.size = Pt(11)
        note.font.italic = True
        note.font.color.rgb = None


def add_inline_runs(paragraph, text):
    # Minimal Markdown inline handling for bold and code.
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            paragraph.add_run(part)


def render():
    doc = Document()
    style_document(doc)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    cover = True
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line == "\\pagebreak":
            doc.add_page_break()
            cover = False
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            rows.pop(1)
            cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=cols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ridx, row in enumerate(rows):
                for cidx in range(cols):
                    set_cell_text(table.cell(ridx, cidx), row[cidx] if cidx < len(row) else "", ridx == 0)
            doc.add_paragraph()
            continue
        image_match = re.match(r"^!\[(.+)\]\((.+)\)$", line)
        if image_match:
            caption, rel_path = image_match.groups()
            image_path = SOURCE.parent / rel_path
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.add_run().add_picture(str(image_path), width=Cm(15.5))
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.first_line_indent = Cm(0)
            cr = cp.add_run(caption)
            cr.bold = True
            cr.italic = True
            cr.font.name = "Times New Roman"
            cr.font.size = Pt(10)
            i += 1
            continue
        if line.startswith("# "):
            if not cover and (line.startswith("# CHƯƠNG ") or line.startswith("# TÀI LIỆU ")):
                doc.add_page_break()
            p = doc.add_paragraph(style="Title" if cover else "Heading 1")
            add_inline_runs(p, line[2:])
            if cover:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Subtitle" if cover else "Heading 2")
            add_inline_runs(p, line[3:])
            if cover:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 3")
            add_inline_runs(p, line[4:])
            i += 1
            continue
        if line.startswith("[MỤC LỤC TỰ ĐỘNG"):
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0)
            add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "Nhấn F9 để cập nhật mục lục")
            i += 1
            continue
        if line.startswith("**[HÌNH"):
            p = doc.add_paragraph(style="Report Note")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline_runs(p, line)
            i += 1
            continue
        if re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.first_line_indent = Cm(0)
            add_inline_runs(p, re.sub(r"^\d+\. ", "", line))
            i += 1
            continue
        p = doc.add_paragraph()
        if cover:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
        add_inline_runs(p, line.replace("  ", " "))
        i += 1

    # Page number in footer.
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, "PAGE", "1")

    # Ask Word to update fields on open.
    settings = doc.settings._element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)

    props = doc.core_properties
    props.title = "Báo cáo chính – ChildGuard – An Open Guard Kids Application"
    props.subject = "Xây dựng ứng dụng hỗ trợ quản lý việc sử dụng máy tính của trẻ em"
    props.author = "Nguyễn Văn Phước; Hà Công Thuận"
    props.comments = "Sinh tự động từ BaoCaoChinh_ChildGuard.md; cần cập nhật các trường trong ngoặc vuông."
    doc.save(OUTPUT)
    # Keep the earlier filename synchronized so existing links do not become stale.
    doc.save(LEGACY_OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    render()
